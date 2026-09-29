import os
from io import BytesIO
from unittest.mock import Mock, patch

os.environ.setdefault("SUPABASE_URL", "https://example.supabase.co")
os.environ.setdefault("SUPABASE_ANON_KEY", "anon-key")

from openpyxl import load_workbook  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from api.main import app  # noqa: E402
from job_radar.models import JobListing  # noqa: E402

client = TestClient(app)


def fake_profiles_response(rows):
    response = Mock()
    response.status_code = 200
    response.json = Mock(return_value=rows)
    response.raise_for_status = Mock()
    return response


def test_generate_requires_authorization_header():
    response = client.post("/generate")
    assert response.status_code == 401


def test_generate_returns_401_when_supabase_rejects_token():
    unauthorized = Mock(status_code=401)
    with patch("api.main.requests.get", return_value=unauthorized):
        response = client.post("/generate", headers={"Authorization": "Bearer bad-token"})
    assert response.status_code == 401


def test_generate_returns_400_when_no_profiles():
    with patch("api.main.requests.get", return_value=fake_profiles_response([])):
        response = client.post("/generate", headers={"Authorization": "Bearer token"})
    assert response.status_code == 400


def test_generate_builds_and_returns_excel():
    rows = [{"id": "p1", "name": "QA Automation", "keywords": {"selenium": 3, "qa": 3}}]
    listing = JobListing(
        title="QA Automation Tester Junior",
        company="Acme",
        url="https://example.com/1",
        description="Selenium, remoto",
        source="fake",
    )

    with patch("api.main.requests.get", return_value=fake_profiles_response(rows)), \
            patch("api.main.ComputrabajoScraper") as mock_cb, \
            patch("api.main.BumeranScraper") as mock_bm:
        mock_cb.return_value.fetch.return_value = [listing]
        mock_bm.return_value.fetch.return_value = []

        response = client.post("/generate", headers={"Authorization": "Bearer token"})

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = load_workbook(BytesIO(response.content))
    assert "QA Automation" in workbook.sheetnames
    sheet = workbook["QA Automation"]
    assert sheet.cell(row=2, column=1).value == "QA Automation Tester Junior"
