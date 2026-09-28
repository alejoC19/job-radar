from openpyxl import load_workbook

from job_radar.config import ProfileConfig
from job_radar.export import build_workbook, save_workbook
from job_radar.models import JobListing, ScoreResult


def make_listing(title: str, url: str) -> JobListing:
    return JobListing(title=title, company="Acme", url=url, source="test")


def make_scored(listing, **scores):
    results = [
        ScoreResult(profile_key=key, profile_name=key, score=score)
        for key, score in scores.items()
    ]
    return listing, results


def test_build_workbook_creates_one_sheet_per_profile():
    profiles = [
        ProfileConfig(key="qa", name="QA Automation"),
        ProfileConfig(key="dev", name="Desarrollador"),
    ]
    listing = make_listing("QA Tester", "https://example.com/1")
    scored = [make_scored(listing, qa=9, dev=2)]

    wb = build_workbook(scored, profiles)
    assert wb.sheetnames == ["QA Automation", "Desarrollador"]


def test_listing_goes_to_sheet_of_highest_scoring_profile():
    profiles = [
        ProfileConfig(key="qa", name="QA Automation"),
        ProfileConfig(key="dev", name="Desarrollador"),
    ]
    listing = make_listing("QA Tester", "https://example.com/1")
    scored = [make_scored(listing, qa=9, dev=2)]

    wb = build_workbook(scored, profiles)
    qa_sheet = wb["QA Automation"]
    dev_sheet = wb["Desarrollador"]

    assert qa_sheet.max_row == 2  # header + 1 row
    assert qa_sheet.cell(row=2, column=1).value == "QA Tester"
    assert dev_sheet.max_row == 1  # solo header, no matches


def test_rows_sorted_by_score_descending():
    profiles = [ProfileConfig(key="qa", name="QA Automation")]
    low = make_listing("Puesto bajo puntaje", "https://example.com/low")
    high = make_listing("Puesto alto puntaje", "https://example.com/high")
    scored = [make_scored(low, qa=3), make_scored(high, qa=9)]

    wb = build_workbook(scored, profiles)
    sheet = wb["QA Automation"]
    assert sheet.cell(row=2, column=1).value == "Puesto alto puntaje"
    assert sheet.cell(row=3, column=1).value == "Puesto bajo puntaje"


def test_save_workbook_writes_readable_file(tmp_path):
    profiles = [ProfileConfig(key="qa", name="QA Automation")]
    listing = make_listing("QA Tester", "https://example.com/1")
    scored = [make_scored(listing, qa=9)]

    wb = build_workbook(scored, profiles)
    out_path = tmp_path / "nested" / "avisos.xlsx"
    save_workbook(wb, out_path)

    assert out_path.exists()
    reloaded = load_workbook(out_path)
    assert reloaded["QA Automation"].cell(row=2, column=1).value == "QA Tester"
